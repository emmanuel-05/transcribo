# app/infrastructure/document_generator.py
import io
import uuid
import datetime
from docxtpl import DocxTemplate
from docx import Document
from botocore.exceptions import ClientError

from app.infrastructure.storage.s3 import s3_client, upload_file_to_s3, generate_presigned_url
from app.core.config import get_settings

settings = get_settings()

BUCKET_NAME = "documents"
TEMPLATE_KEY = "templates/lettre_standard.docx"

async def ensure_default_template_exists():
    """Vérifie que le template par défaut existe dans S3, sinon le crée."""
    try:
        s3_client.head_bucket(Bucket=BUCKET_NAME)
    except ClientError:
        s3_client.create_bucket(Bucket=BUCKET_NAME)

    try:
        s3_client.head_object(Bucket=BUCKET_NAME, Key=TEMPLATE_KEY)
    except ClientError as e:
        if e.response['Error']['Code'] == "404":
            # Le template n'existe pas, on le crée en mémoire
            doc = Document()
            doc.add_paragraph("Expéditeur : {{ expediteur }}")
            doc.add_paragraph("Destinataire : {{ destinataire }}")
            doc.add_paragraph("Date : {{ date }}")
            doc.add_paragraph("Objet : {{ objet }}")
            doc.add_paragraph("")
            doc.add_paragraph("{{ contenu }}")
            
            # Sauvegarder en bytes
            file_stream = io.BytesIO()
            doc.save(file_stream)
            file_stream.seek(0)
            
            # Upload sur S3
            s3_client.put_object(
                Bucket=BUCKET_NAME,
                Key=TEMPLATE_KEY,
                Body=file_stream.read(),
                ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        else:
            raise

async def generate_docx_from_transcript(project_id: str, transcript_id: str, text_content: str, expediteur: str, destinataire: str, objet: str) -> str:
    """
    Génère un document Word (.docx) à partir d'un texte et de variables.
    Retourne l'URL présignée S3 pour téléchargement.
    """
    await ensure_default_template_exists()
    
    # 1. Télécharger le template depuis S3
    response = s3_client.get_object(Bucket=BUCKET_NAME, Key=TEMPLATE_KEY)
    template_bytes = response['Body'].read()
    
    # 2. Charger avec DocxTemplate
    template_stream = io.BytesIO(template_bytes)
    doc = DocxTemplate(template_stream)
    
    # 3. Remplir les variables
    context = {
        'expediteur': expediteur,
        'destinataire': destinataire,
        'date': datetime.datetime.now().strftime("%d/%m/%Y"),
        'objet': objet,
        'contenu': text_content
    }
    doc.render(context)
    
    # 4. Sauvegarder le résultat en mémoire
    output_stream = io.BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    
    # 5. Upload sur S3
    doc_id = str(uuid.uuid4())
    doc_key = f"projects/{project_id}/documents/{doc_id}.docx"
    
    await upload_file_to_s3(
        file_data=output_stream.read(),
        bucket=BUCKET_NAME,
        key=doc_key,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    
    # 6. Retourner l'URL présignée
    url = generate_presigned_url(bucket=BUCKET_NAME, key=doc_key)
    return url
